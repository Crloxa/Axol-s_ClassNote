"""端到端冒烟测试（不联网调用模型）。

前置：后端已在 127.0.0.1:18471 运行；用法：
    .venv/Scripts/python backend/tests/smoke_api.py [speech.wav]
传入 wav 时会额外验证录音/转写链路；wav 必须位于项目目录内（只读该文件）。
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import httpx

BASE = "http://127.0.0.1:18471/api"
PROJECT_DIR = Path(__file__).resolve().parents[2]
client = httpx.Client(base_url=BASE, timeout=120, trust_env=False)  # 本机测试不走系统代理

passed: list[str] = []


def ok(name: str, cond: bool, extra: str = ""):
    status = "PASS" if cond else "FAIL"
    passed.append(f"[{status}] {name}{(' - ' + extra) if extra else ''}")
    print(passed[-1])
    if not cond:
        sys.exit(1)


def safe_wav_path(arg: str) -> Path | None:
    """规范化命令行传入的 wav 路径：必须是项目目录内的已存在 .wav 文件。"""
    p = Path(arg).resolve()
    if p.suffix.lower() != ".wav" or not p.is_file():
        return None
    if PROJECT_DIR != p and PROJECT_DIR not in p.parents:
        return None
    return p


def make_pptx() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    for title, body in [
        ("梯度下降", "目标：最小化损失函数\n沿负梯度方向更新参数"),
        ("学习率", "步长过大：震荡或不收敛\n步长过小：收敛太慢"),
        ("小结", "梯度下降是优化的基础工具"),
    ]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("参考讲义", level=1)
    d.add_heading("第一节 梯度", level=2)
    d.add_paragraph("梯度是函数上升最快的方向，负梯度即下降最快的方向。")
    d.add_heading("第二节 常见问题", level=2)
    d.add_paragraph("学习率需要根据损失曲面调整。")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pdf() -> bytes:
    # 手工构造一个包含可提取文本的最小 PDF（ASCII）
    text = "(Gradient descent review sheet. Page one.)"
    content = f"BT /F1 18 Tf 72 700 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref_pos = out.tell()
    out.write(f"xref\n0 {len(objects)+1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF".encode())
    return out.getvalue()


def main():
    r = client.get("/lessons")
    ok("后端可达 GET /lessons", r.status_code == 200)

    r = client.post("/lessons", json={"name": "冒烟测试课堂"})
    ok("创建课堂", r.status_code == 200)
    lid = r.json()["id"]

    r = client.post(f"/lessons/{lid}/ppt",
                    files={"file": ("lec01.pptx", make_pptx())})
    ok("导入主 PPTX", r.status_code == 200 and r.json()["slide_count"] == 3,
       f"{r.json().get('slide_count', '?')} 页")

    for name, blob in [("讲义.docx", make_docx()), ("notes.md", "## 附注\n梯度噪声很重要".encode()),
                       ("sheet.pdf", make_pdf())]:
        r = client.post(f"/lessons/{lid}/sources", files={"file": (name, blob)})
        ok(f"导入参考资料 {name}", r.status_code == 200)

    r = client.get(f"/lessons/{lid}")
    data = r.json()
    ok("课堂详情：3 页 PPT", len(data["slides"]) == 3)
    ok("课堂详情：3 份资料", len(data["sources"]) == 3)
    pdf_page1 = data["sources"][2]["content"]["pages"][0]
    ok("PDF 文本已提取", not pdf_page1["empty"], pdf_page1["text"][:30])

    # ---- 本地规则生成（详细提纲）----
    src_ids = [s["id"] for s in data["sources"]]
    r = client.post("/generate/preview", json={
        "lesson_id": lid, "mode": "outline", "engine": "local",
        "slide_idxs": None, "source_ids": src_ids, "instruction": ""})
    ok("本地预览", r.status_code == 200, f"{r.json()['stats']['total_chars']} 字符")
    token = r.json()["token"]

    r = client.post("/generate", json={"token": token})
    body = r.json()["content_md"]
    ok("本地生成详细提纲", r.status_code == 200 and "[PPT 第 1 页]" in body)
    ok("提纲含参考资料位置", "[sheet.pdf：第 1 页]" in body or "[notes.md" in body)

    # 令牌一次性
    r2 = client.post("/generate", json={"token": token})
    ok("一次性令牌已失效", r2.status_code == 400)

    # ---- PPT 外补充模式（本地，单章）----
    r = client.post("/generate/preview", json={
        "lesson_id": lid, "mode": "supplement", "engine": "local",
        "slide_idxs": [1], "source_ids": []})
    r2 = client.post("/generate", json={"token": r.json()["token"]})
    ok("单章本地生成（PPT 外补充）", r2.status_code == 200 and "PPT 外补充" in r2.json()["content_md"])

    # ---- 预览（模型引擎）不联网，Key 未保存 ----
    r = client.post("/generate/preview", json={
        "lesson_id": lid, "mode": "key_points", "engine": "model",
        "provider": "deepseek", "slide_idxs": None, "source_ids": []})
    pv = r.json()
    ok("远程预览显示服务商信息", pv["provider"]["label"] == "DeepSeek" and pv["provider"]["has_key"] is False)

    # ---- 出站安全：拒绝内网端点 ----
    r = client.put("/providers/deepseek", json={"base_url": "http://127.0.0.1:18471/api"})
    ok("SSRF 防护拒绝内网端点", r.status_code == 400, (r.json().get("detail") or "")[:40])

    # ---- 录音转写链路（可选，传 wav 才测）----
    wav = safe_wav_path(sys.argv[1]) if len(sys.argv) > 1 else None
    if wav is not None:
        client.put("/settings", json={"stt_model": "tiny", "stt_language": ""})
        client.post(f"/lessons/{lid}/recording/start")
        r = client.post(f"/lessons/{lid}/audio", content=wav.read_bytes())
        ok("上传音频分块", r.status_code == 200)
        r = client.post(f"/lessons/{lid}/transcribe/flush")
        segs = r.json()["segments"]
        ok("VAD 切分并转写出片段", len(segs) > 0, f"{len(segs)} 段")
        print("    转写内容:", [s["text"] for s in segs])
        # flush 返回值不含 id，从详情接口取最新的带 id 片段
        all_segs = client.get(f"/lessons/{lid}").json()["segments"]
        segs = all_segs[-len(segs):] if segs else []
        if segs:
            seg_id = segs[0]["id"]
            r = client.patch(f"/segments/{seg_id}", json={"highlight": True})
            ok("片段标重点", r.json()["highlight"] == 1)
            r = client.patch(f"/segments/{seg_id}", json={"manual_slide": 2})
            ok("片段关联页码", r.json()["manual_slide"] == 2)
            r = client.patch(f"/segments/{seg_id}", json={"text": segs[0]["text"] + "（已修正）"})
            ok("片段手动修正", "已修正" in r.json()["text"])
        r = client.post(f"/lessons/{lid}/recording/stop")
        ok("停止录音", r.status_code == 200)
        client.put("/settings", json={"stt_language": "zh"})
        r = client.post("/generate/preview", json={
            "lesson_id": lid, "mode": "key_points", "engine": "local",
            "slide_idxs": None, "source_ids": []})
        r2 = client.post("/generate", json={"token": r.json()["token"]})
        ok("重点讲义包含转写", "课堂转写" in r2.json()["content_md"])
    else:
        print("[SKIP] 录音转写链路（未提供项目目录内的 wav）")

    # ---- 导出 ----
    r = client.get("/lessons/" + lid)
    rid = r.json()["results"][0]["id"]
    r = client.get(f"/results/{rid}/export", params={"format": "md"})
    ok("导出 Markdown", r.status_code == 200 and len(r.content) > 100)
    r = client.get(f"/results/{rid}/export", params={"format": "html"})
    ok("导出 HTML", r.status_code == 200 and b"<html" in r.content[:200])

    print(f"\n全部通过：{sum(1 for p in passed if p.startswith('[PASS]'))} 项")


if __name__ == "__main__":
    main()
