"""本地文件解析：只提取文本与来源位置，不做 OCR。

- PPTX  → 每页 {idx, title, text}，页码从 1 开始展示
- PDF   → 每页 {page, text, empty}；扫描页 empty=True，提示“未提取文字”
- DOCX  → 按文档顺序的块 {type: heading|para|table, level, text, path}
- MD    → 按标题切分的 {level, title, text, path}
- 代码  → UTF-8 全文 + 行数，来源定位用行号
"""
from __future__ import annotations

import io
import json
from pathlib import Path

CODE_EXTS = {
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".go", ".rs",
    ".c", ".cpp", ".h", ".cs", ".sql", ".json", ".yaml", ".yml",
    ".toml", ".sh",
}


def parse_pptx(data: bytes) -> list[dict]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slides: list[dict] = []
    for i, slide in enumerate(prs.slides):
        title = ""
        texts: list[str] = []
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text_frame.text.strip()
        except Exception:
            title = ""
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                line = "".join(run.text for run in para.runs).strip()
                if not line:
                    continue
                if shape == slide.shapes.title:
                    continue
                texts.append(line)
        # 有些 PPT 标题占位符不在 shapes.title 里，用首个非空文本兜底
        if not title and texts:
            title = texts.pop(0)
        slides.append({"idx": i + 1, "title": title or "(无标题)", "text": "\n".join(texts)})
    return slides


def parse_pdf(data: bytes) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:
            text = ""
        pages.append({"page": i + 1, "text": text, "empty": not text})
    return {"pages": pages}


def _docx_style_level(para) -> int | None:
    name = (para.style.name or "").lower() if para.style is not None else ""
    for i in range(1, 7):
        if name == f"heading {i}" or name == f"标题 {i}":
            return i
    return None


def parse_docx(data: bytes) -> dict:
    import docx
    from docx.document import Document as _Doc
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(io.BytesIO(data))

    def iter_blocks(parent: _Doc):
        for child in parent.element.body.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                yield Table(child, parent)

    blocks: list[dict] = []
    stack: list[tuple[int, str]] = []  # [(level, title)]

    def path_of(level: int, title: str) -> str:
        keep = [t for lv, t in stack if lv < level]
        return " > ".join(keep + [title])

    for block in iter_blocks(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            level = _docx_style_level(block)
            if level:
                stack[:] = [(lv, t) for lv, t in stack if lv < level]
                stack.append((level, text))
                blocks.append({"type": "heading", "level": level, "text": text,
                               "path": path_of(level, text)})
            else:
                blocks.append({"type": "para", "level": 0, "text": text,
                               "path": path_of(99, "") if stack else ""})
        else:  # Table
            rows = []
            for row in block.rows:
                cells = [c.text.strip() for c in row.cells]
                rows.append(" | ".join(cells))
            blocks.append({"type": "table", "level": 0, "text": "\n".join(rows),
                           "path": path_of(99, "") if stack else ""})
    return {"blocks": blocks}


def parse_markdown_text(name: str, text: str) -> dict:
    sections: list[dict] = []
    stack: list[tuple[int, str]] = []
    current: dict | None = None
    lines = text.splitlines()
    for line in lines:
        stripped = line.strip()
        level = 0
        if stripped.startswith("#"):
            hashes = len(stripped) - len(stripped.lstrip("#"))
            rest = stripped[hashes:].strip()
            if 1 <= hashes <= 6 and (not rest or rest.startswith(" ") or rest.startswith("#")):
                level = hashes
                title = stripped[hashes:].strip().lstrip("#").strip()
                stack[:] = [(lv, t) for lv, t in stack if lv < level]
                stack.append((level, title))
                current = {"level": level, "title": title,
                           "path": " > ".join(t for _, t in stack), "text": ""}
                sections.append(current)
                continue
        if current is None:
            current = {"level": 0, "title": "(正文开头)", "path": "", "text": ""}
            sections.append(current)
        current["text"] += line + "\n"
    for s in sections:
        s["text"] = s["text"].strip()
    return {"name": name, "sections": sections}


def parse_code(data: bytes) -> dict:
    text = data.decode("utf-8", errors="replace")
    return {"text": text, "lines": text.count("\n") + (1 if text and not text.endswith("\n") else 0)}


def kind_for(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pptx":
        return "pptx"
    if ext == ".pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"
    if ext in (".md", ".markdown"):
        return "md"
    if ext in CODE_EXTS:
        return "code"
    return ""


def parse_source(name: str, kind: str, data: bytes) -> dict:
    """返回可 JSON 序列化的解析结果（存 sources.content）。"""
    if kind == "pdf":
        return parse_pdf(data)
    if kind == "docx":
        return parse_docx(data)
    if kind == "md":
        return parse_markdown_text(name, data.decode("utf-8", errors="replace"))
    if kind == "code":
        return parse_code(data)
    raise ValueError(f"不支持的资料类型: {kind}")


def source_items(name: str, kind: str, content: dict) -> list[dict]:
    """把解析结果展开为 [(来源位置标签, 文本)] 列表，用于预览/生成。"""
    items: list[dict] = []
    if kind == "pdf":
        for p in content.get("pages", []):
            label = f"第 {p['page']} 页" + ("（未提取文字，可能是扫描页）" if p["empty"] else "")
            items.append({"loc": label, "text": p["text"]})
    elif kind == "docx":
        for b in content.get("blocks", []):
            loc = b.get("path") or "正文"
            items.append({"loc": loc, "text": b["text"]})
    elif kind == "md":
        for s in content.get("sections", []):
            loc = s.get("path") or "正文"
            items.append({"loc": loc, "text": s["text"]})
    elif kind == "code":
        text = content.get("text", "")
        lines = text.splitlines()
        step = 80
        for i in range(0, len(lines), step):
            chunk = "\n".join(lines[i:i + step])
            if chunk.strip():
                items.append({"loc": f"第 {i + 1}-{min(i + step, len(lines))} 行", "text": chunk})
        if not items:
            items.append({"loc": "全文", "text": ""})
    return items


def dumps(obj) -> str:
    return json.dumps(obj, ensure_ascii=False)
